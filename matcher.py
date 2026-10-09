import re
from pathlib import Path
from collections import Counter

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

SKILL_KEYWORDS = [
    "python", "java", "c++", "c#", "javascript", "typescript", "html", "css",
    "react", "node.js", "flask", "django", "fastapi", "sql", "mysql",
    "postgresql", "mongodb", "sqlite", "machine learning", "deep learning",
    "artificial intelligence", "nlp", "natural language processing",
    "computer vision", "data analysis", "data science", "pandas", "numpy",
    "scikit-learn", "tensorflow", "pytorch", "excel", "power bi", "tableau",
    "aws", "azure", "google cloud", "docker", "kubernetes", "git", "github",
    "linux", "rest api", "api", "statistics", "communication", "leadership",
    "project management", "problem solving", "cybersecurity", "devops",
    "sentence-transformers", "spacy", "scikit learn", "data structures",
    "oop", "object oriented programming"
]

def extract_text_from_file(path):
    """Extract text from PDF, DOCX, or TXT. Scanned PDFs need OCR."""
    path = Path(path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        import fitz  # PyMuPDF
        document = fitz.open(path)
        return "\n".join(page.get_text("text") for page in document)

    if ext == ".docx":
        from docx import Document
        document = Document(path)
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)

    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")

    raise ValueError("Unsupported file type")

def clean_text(text):
    text = text.lower()
    text = re.sub(r"[^a-z0-9+#./\s-]", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def extract_skills(text):
    """Simple keyword-based skill extraction for the student prototype."""
    cleaned = clean_text(text)
    found = []
    for skill in SKILL_KEYWORDS:
        pattern = r"(?<![a-z0-9])" + re.escape(skill.lower()) + r"(?![a-z0-9])"
        if re.search(pattern, cleaned):
            # Normalize a couple of equivalent labels for display.
            display = "scikit-learn" if skill == "scikit learn" else skill
            if display not in found:
                found.append(display)
    return sorted(found)

def analyse_resume(resume_text, job_description):
    resume_clean = clean_text(resume_text)
    job_clean = clean_text(job_description)

    resume_skills = set(extract_skills(resume_text))
    job_skills = set(extract_skills(job_description))
    matched = sorted(resume_skills & job_skills)
    missing = sorted(job_skills - resume_skills)

    # Semantic-like text similarity baseline using TF-IDF + cosine similarity.
    # It is a baseline, not a validated hiring prediction.
    try:
        vectorizer = TfidfVectorizer(ngram_range=(1, 2), stop_words="english")
        matrix = vectorizer.fit_transform([resume_clean, job_clean])
        similarity = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0])
    except ValueError:
        similarity = 0.0

    # Combined score: 70% text similarity and 30% skill overlap.
    # If no skills were found in the job description, use similarity only.
    if job_skills:
        skill_score = len(matched) / len(job_skills)
        score = round(100 * (0.70 * similarity + 0.30 * skill_score))
    else:
        score = round(100 * similarity)

    score = max(0, min(100, score))
    return {
        "score": score,
        "similarity": round(similarity * 100),
        "matched_skills": matched,
        "missing_skills": missing,
        "resume_skills": sorted(resume_skills),
        "job_skills": sorted(job_skills),
        "explanation": (
            "Score combines TF-IDF text similarity and overlap of detected skills. "
            "It is a prototype indicator, not a hiring decision."
        )
    }
